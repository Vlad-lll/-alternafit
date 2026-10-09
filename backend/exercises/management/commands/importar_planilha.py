"""
Importa a planilha do banco de exercícios (formato "modelo de banco de dados",
versão 2.2 em diante: uma aba por tabela, com IDs fixos) e SUBSTITUI todo o
catálogo de exercícios do banco pelo conteúdo da planilha.

Uso (dentro da pasta backend/):
    venv\\Scripts\\python.exe manage.py importar_planilha "..\\data\\arquivo.xlsx" --dry-run
    venv\\Scripts\\python.exe manage.py importar_planilha "..\\data\\arquivo.xlsx"

- --dry-run simula tudo e mostra o relatório, mas não grava nada.
- Antes de gravar, a planilha inteira é conferida (IDs, referências entre
  abas, duplicatas). Se houver qualquer erro, nada é gravado.
- Tudo acontece dentro de uma transação: se der erro no meio, o banco volta
  exatamente como estava.
- Os IDs da planilha viram os IDs do banco, então o "codigo" de cada
  exercício na API continua o mesmo entre uma importação e outra.
- Usuários e logins do Admin não são afetados.
"""

import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError
from django.core.validators import URLValidator
from django.db import transaction
from openpyxl import load_workbook

from exercises.models import (
    Difficulty,
    Equipment,
    Exercise,
    ExerciseAlias,
    ExerciseAlternative,
    ExerciseMuscularGroup,
    ExerciseStep,
    MovementPattern,
    MuscularGroup,
    Source,
)

# Colunas obrigatórias de cada aba. Colunas extras (como _aux_ e _calc_) são ignoradas.
SHEETS = {
    'nivel_dificuldade': ['id', 'nome', 'descricao'],
    'grupo_muscular': ['id', 'codigo', 'nome', 'regiao', 'descricao'],
    'equipamento': ['id', 'nome', 'categoria', 'descricao'],
    'padrao_movimento': ['id', 'codigo', 'nome'],
    'fonte': ['id', 'autor_organizacao', 'titulo', 'url', 'tipo'],
    'exercicio': [
        'id', 'nome', 'nome_en', 'descricao', 'nivel_id', 'padrao_movimento_id',
        'dica_execucao', 'video_url', 'status_dados', 'observacao',
    ],
    'exercicio_grupo_muscular': ['exercicio_id', 'grupo_muscular_id', 'papel'],
    'exercicio_equipamento': ['exercicio_id', 'equipamento_id'],
    'exercicio_alternativo': ['exercicio_id', 'alternativo_id', 'ordem', 'observacao'],
    'exercicio_passo': ['exercicio_id', 'ordem', 'descricao', 'dica_extra'],
    'exercicio_apelido': ['exercicio_id', 'apelido'],
    'exercicio_fonte': ['exercicio_id', 'fonte_id'],
}

ROLE_VALUES = {
    'primario': ExerciseMuscularGroup.Role.PRIMARY,
    'principal': ExerciseMuscularGroup.Role.PRIMARY,
    'secundario': ExerciseMuscularGroup.Role.SECONDARY,
    'auxiliar': ExerciseMuscularGroup.Role.SECONDARY,
}


def text(value):
    if value is None:
        return ''
    return str(value).strip()


def fold(value):
    """Minúsculas, sem acentos e sem espaços repetidos, para comparar textos."""
    normalized = unicodedata.normalize('NFKD', text(value).lower())
    normalized = ''.join(ch for ch in normalized if not unicodedata.combining(ch))
    return re.sub(r'\s+', ' ', normalized)


def to_int(value):
    try:
        number = float(text(value).replace(',', '.'))
    except ValueError:
        return None
    return int(number) if number.is_integer() else None


def choice_lookup(choices):
    """Permite achar o valor de um TextChoices pelo texto que aparece na planilha (o rótulo)."""
    return {fold(label): value for value, label in choices}


def is_valid_url(url):
    try:
        URLValidator()(url)
    except ValidationError:
        return False
    return True


class Command(BaseCommand):
    help = 'Substitui o catálogo de exercícios do banco pelo conteúdo da planilha (formato v2.2 em diante).'

    def add_arguments(self, parser):
        parser.add_argument('arquivo', help='Caminho do arquivo .xlsx')
        parser.add_argument(
            '--dry-run',
            action='store_true',
            help='Simula a importação e mostra o relatório, sem gravar nada no banco.',
        )
        parser.add_argument(
            '--detalhes',
            action='store_true',
            help='Lista também as trocas "fracas" (alternativa sem o mesmo grupo principal nem o mesmo padrão).',
        )

    def handle(self, *args, **options):
        path = Path(options['arquivo'])
        if not path.exists():
            raise CommandError(f'Arquivo não encontrado: {path}')

        self.errors = []
        self.warnings = []
        tables = self.read_workbook(path)
        self.validate(tables)
        if self.errors:
            self.stdout.write(self.style.ERROR(f'\nA planilha tem {len(self.errors)} erro(s). Nada foi gravado:'))
            for error in self.errors:
                self.stdout.write(self.style.ERROR(f'  - {error}'))
            raise CommandError('Corrija a planilha e rode de novo.')

        with transaction.atomic():
            removed = self.clear_catalog()
            created = self.write(tables)
            if options['dry_run']:
                transaction.set_rollback(True)

        self.print_report(options, removed, created, self.check_rules(tables))

    # ------------------------------------------------------------------ leitura

    def read_workbook(self, path):
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            tables = {}
            for sheet_name, required in SHEETS.items():
                if sheet_name not in workbook.sheetnames:
                    raise CommandError(f'A planilha não tem a aba "{sheet_name}". Ela está no formato v2.2 ou mais novo?')
                rows = workbook[sheet_name].iter_rows(values_only=True)
                header = [text(cell) for cell in next(rows, ())]
                missing = [column for column in required if column not in header]
                if missing:
                    raise CommandError(f'Aba "{sheet_name}": faltam as colunas {missing}.')
                records = []
                for number, values in enumerate(rows, start=2):
                    record = {column: values[header.index(column)] if header.index(column) < len(values) else None
                              for column in required}
                    if all(text(value) == '' for value in record.values()):
                        continue
                    record['_linha'] = number
                    records.append(record)
                tables[sheet_name] = records
            return tables
        finally:
            workbook.close()

    # --------------------------------------------------------------- conferência

    def error(self, sheet, record, message):
        self.errors.append(f'{sheet}, linha {record["_linha"]}: {message}')

    def validate(self, tables):
        ids = {}
        for sheet in ['nivel_dificuldade', 'grupo_muscular', 'equipamento', 'padrao_movimento', 'fonte', 'exercicio']:
            seen = set()
            for record in tables[sheet]:
                record['id'] = to_int(record['id'])
                if record['id'] is None or record['id'] < 1:
                    self.error(sheet, record, 'id vazio ou inválido.')
                elif record['id'] in seen:
                    self.error(sheet, record, f'id {record["id"]} repetido.')
                seen.add(record['id'])
            ids[sheet] = seen

        for record in tables['nivel_dificuldade']:
            if record['id'] not in (1, 2, 3):
                self.error('nivel_dificuldade', record, 'o id do nível precisa ser 1, 2 ou 3.')

        self.check_unique_text(tables, 'grupo_muscular', 'nome')
        self.check_unique_text(tables, 'equipamento', 'nome')
        self.check_unique_text(tables, 'padrao_movimento', 'codigo')
        self.check_unique_text(tables, 'padrao_movimento', 'nome')
        self.check_unique_text(tables, 'exercicio', 'nome')
        self.check_unique_text(tables, 'exercicio_apelido', 'apelido')

        for record in tables['fonte']:
            if not text(record['titulo']) or not text(record['autor_organizacao']):
                self.error('fonte', record, 'autor e título são obrigatórios.')
            if not is_valid_url(text(record['url'])):
                self.error('fonte', record, f'URL inválida: "{text(record["url"])}".')

        for record in tables['exercicio']:
            for column, sheet in [('nivel_id', 'nivel_dificuldade'), ('padrao_movimento_id', 'padrao_movimento')]:
                if text(record[column]):
                    record[column] = to_int(record[column])
                    if record[column] not in ids[sheet]:
                        self.error('exercicio', record, f'{column} {record[column]} não existe na aba {sheet}.')
                else:
                    record[column] = None
            video = text(record['video_url'])
            if video and not is_valid_url(video):
                self.warnings.append(f'exercicio, linha {record["_linha"]}: link de vídeo inválido. Ignorado.')
                record['video_url'] = ''

        links = [
            ('exercicio_grupo_muscular', [('exercicio_id', 'exercicio'), ('grupo_muscular_id', 'grupo_muscular')]),
            ('exercicio_equipamento', [('exercicio_id', 'exercicio'), ('equipamento_id', 'equipamento')]),
            ('exercicio_alternativo', [('exercicio_id', 'exercicio'), ('alternativo_id', 'exercicio')]),
            ('exercicio_fonte', [('exercicio_id', 'exercicio'), ('fonte_id', 'fonte')]),
            ('exercicio_passo', [('exercicio_id', 'exercicio')]),
            ('exercicio_apelido', [('exercicio_id', 'exercicio')]),
        ]
        for sheet, references in links:
            for record in tables[sheet]:
                for column, target in references:
                    record[column] = to_int(record[column])
                    if record[column] not in ids[target]:
                        self.error(sheet, record, f'{column} {record[column]} não existe na aba {target}.')

        self.check_unique_pair(tables, 'exercicio_grupo_muscular', 'exercicio_id', 'grupo_muscular_id')
        self.check_unique_pair(tables, 'exercicio_equipamento', 'exercicio_id', 'equipamento_id')
        self.check_unique_pair(tables, 'exercicio_alternativo', 'exercicio_id', 'alternativo_id')
        self.check_unique_pair(tables, 'exercicio_fonte', 'exercicio_id', 'fonte_id')
        self.check_unique_pair(tables, 'exercicio_passo', 'exercicio_id', 'ordem')

        for record in tables['exercicio_grupo_muscular']:
            role = ROLE_VALUES.get(fold(record['papel']))
            if role is None:
                self.error('exercicio_grupo_muscular', record, f'papel "{text(record["papel"])}" inválido (use PRIMARIO ou SECUNDARIO).')
            record['papel'] = role

        for sheet in ['exercicio_alternativo', 'exercicio_passo']:
            for record in tables[sheet]:
                record['ordem'] = to_int(record['ordem'])
                if record['ordem'] is None or record['ordem'] < 1:
                    self.error(sheet, record, 'ordem vazia ou inválida.')

        for record in tables['exercicio_alternativo']:
            if record['exercicio_id'] == record['alternativo_id']:
                self.error('exercicio_alternativo', record, 'um exercício não pode ser alternativa de si mesmo.')

        exercise_names = {fold(r['nome']): r['id'] for r in tables['exercicio']}
        for record in tables['exercicio_apelido']:
            owner = exercise_names.get(fold(record['apelido']))
            if owner is not None and owner != record['exercicio_id']:
                self.error('exercicio_apelido', record, f'o apelido "{text(record["apelido"])}" é o nome de outro exercício.')
        for record in tables['exercicio_passo']:
            if not text(record['descricao']):
                self.error('exercicio_passo', record, 'descrição do passo vazia.')

    def check_unique_text(self, tables, sheet, column):
        seen = {}
        for record in tables[sheet]:
            key = fold(record[column])
            if not key:
                self.error(sheet, record, f'{column} vazio.')
            elif key in seen:
                self.error(sheet, record, f'{column} "{text(record[column])}" repetido (já aparece na linha {seen[key]}).')
            else:
                seen[key] = record['_linha']

    def check_unique_pair(self, tables, sheet, first, second):
        seen = {}
        for record in tables[sheet]:
            key = (record[first], to_int(record[second]))
            if key in seen:
                self.error(sheet, record, f'combinação repetida (já aparece na linha {seen[key]}).')
            seen[key] = record['_linha']

    # ------------------------------------------------------------------ gravação

    def clear_catalog(self):
        """Apaga o catálogo atual (na ordem certa, por causa das referências). Usuários não são afetados."""
        removed = {'Exercícios': Exercise.objects.count()}
        for model in [
            ExerciseAlternative, ExerciseStep, ExerciseAlias, ExerciseMuscularGroup, Exercise,
            Source, MovementPattern, Equipment, MuscularGroup, Difficulty,
        ]:
            model.objects.all().delete()
        return removed

    def write(self, tables):
        regions = choice_lookup(MuscularGroup.Region.choices)
        categories = choice_lookup(Equipment.Category.choices)
        statuses = choice_lookup(Exercise.DataStatus.choices)

        Difficulty.objects.bulk_create(
            Difficulty(id=r['id'], level=r['id'], name=text(r['nome']), description=text(r['descricao']))
            for r in tables['nivel_dificuldade']
        )
        groups = []
        for r in tables['grupo_muscular']:
            region = regions.get(fold(r['regiao']), '')
            if text(r['regiao']) and not region:
                self.warnings.append(f'grupo_muscular, linha {r["_linha"]}: região "{text(r["regiao"])}" desconhecida. Deixada em branco.')
            groups.append(MuscularGroup(
                id=r['id'], code=text(r['codigo']), name=text(r['nome']), region=region, description=text(r['descricao']),
            ))
        MuscularGroup.objects.bulk_create(groups)
        equipment = []
        for r in tables['equipamento']:
            category = categories.get(fold(r['categoria']), '')
            if text(r['categoria']) and not category:
                self.warnings.append(f'equipamento, linha {r["_linha"]}: categoria "{text(r["categoria"])}" desconhecida. Deixada em branco.')
            equipment.append(Equipment(id=r['id'], name=text(r['nome']), category=category, description=text(r['descricao'])))
        Equipment.objects.bulk_create(equipment)
        MovementPattern.objects.bulk_create(
            MovementPattern(id=r['id'], code=text(r['codigo']), name=text(r['nome'])) for r in tables['padrao_movimento']
        )
        Source.objects.bulk_create(
            Source(
                id=r['id'], author=text(r['autor_organizacao']), title=text(r['titulo']),
                url=text(r['url']), source_type=text(r['tipo']),
            )
            for r in tables['fonte']
        )

        exercises = []
        for r in tables['exercicio']:
            status = statuses.get(fold(r['status_dados']), '')
            if text(r['status_dados']) and not status:
                self.warnings.append(f'exercicio, linha {r["_linha"]}: status "{text(r["status_dados"])}" desconhecido. Deixado em branco.')
            exercises.append(Exercise(
                id=r['id'], name=text(r['nome']), name_en=text(r['nome_en']), description=text(r['descricao']),
                tip=text(r['dica_execucao']), video_url=text(r['video_url']), difficulty_id=r['nivel_id'],
                movement_pattern_id=r['padrao_movimento_id'], data_status=status, notes=text(r['observacao']),
            ))
        Exercise.objects.bulk_create(exercises)

        ExerciseMuscularGroup.objects.bulk_create(
            ExerciseMuscularGroup(exercise_id=r['exercicio_id'], muscular_group_id=r['grupo_muscular_id'], role=r['papel'])
            for r in tables['exercicio_grupo_muscular']
        )
        Exercise.equipment.through.objects.bulk_create(
            Exercise.equipment.through(exercise_id=r['exercicio_id'], equipment_id=r['equipamento_id'])
            for r in tables['exercicio_equipamento']
        )
        Exercise.sources.through.objects.bulk_create(
            Exercise.sources.through(exercise_id=r['exercicio_id'], source_id=r['fonte_id'])
            for r in tables['exercicio_fonte']
        )
        ExerciseAlternative.objects.bulk_create(
            ExerciseAlternative(
                exercise_id=r['exercicio_id'], alternative_id=r['alternativo_id'],
                order=r['ordem'], note=text(r['observacao']),
            )
            for r in tables['exercicio_alternativo']
        )
        ExerciseStep.objects.bulk_create(
            ExerciseStep(
                exercise_id=r['exercicio_id'], order=r['ordem'],
                description=text(r['descricao']), extra_tip=text(r['dica_extra']),
            )
            for r in tables['exercicio_passo']
        )
        ExerciseAlias.objects.bulk_create(
            ExerciseAlias(exercise_id=r['exercicio_id'], alias=text(r['apelido'])) for r in tables['exercicio_apelido']
        )

        return {
            'Dificuldades': Difficulty.objects.count(),
            'Grupos musculares': MuscularGroup.objects.count(),
            'Equipamentos': Equipment.objects.count(),
            'Padrões de movimento': MovementPattern.objects.count(),
            'Fontes': Source.objects.count(),
            'Exercícios': Exercise.objects.count(),
            'Ligações exercício-grupo': ExerciseMuscularGroup.objects.count(),
            'Alternativas': ExerciseAlternative.objects.count(),
            'Passos de execução': ExerciseStep.objects.count(),
            'Apelidos': ExerciseAlias.objects.count(),
        }

    # ------------------------------------------------------------- regras do app

    def check_rules(self, tables):
        names = {r['id']: text(r['nome']) for r in tables['exercicio']}
        pattern = {r['id']: r['padrao_movimento_id'] for r in tables['exercicio']}
        with_equipment = {r['exercicio_id'] for r in tables['exercicio_equipamento']}
        primary = defaultdict(set)
        for r in tables['exercicio_grupo_muscular']:
            if r['papel'] == ExerciseMuscularGroup.Role.PRIMARY:
                primary[r['exercicio_id']].add(r['grupo_muscular_id'])
        alternatives = defaultdict(list)
        for r in tables['exercicio_alternativo']:
            alternatives[r['exercicio_id']].append(r['alternativo_id'])

        problems, weak = [], []
        for exercise_id, name in names.items():
            options = alternatives[exercise_id]
            if not 2 <= len(options) <= 4:
                problems.append(f'"{name}" tem {len(options)} alternativa(s); a regra é de 2 a 4.')
            if options and all(option in with_equipment for option in options):
                problems.append(f'"{name}" não tem nenhuma alternativa sem equipamento.')
            if not primary[exercise_id]:
                problems.append(f'"{name}" não tem grupo muscular principal.')
            for option in options:
                if not (primary[exercise_id] & primary[option]) and pattern[exercise_id] != pattern[option]:
                    weak.append(f'{name} → {names[option]}')

        review = Counter(text(r['status_dados']) for r in tables['exercicio'])
        flagged_sources = [text(r['titulo']) for r in tables['fonte'] if 'reconferir' in fold(r['tipo'])]
        return {
            'problems': problems,
            'weak': weak,
            'review': {status: n for status, n in review.items() if fold(status) != 'com fonte'},
            'flagged_sources': flagged_sources,
            'free': len(names) - len(with_equipment & set(names)),
        }

    # ---------------------------------------------------------------- relatório

    def print_report(self, options, removed, created, rules):
        out = self.stdout
        out.write('')
        out.write('========== RELATÓRIO DA IMPORTAÇÃO ==========')
        if options['dry_run']:
            out.write(self.style.WARNING('SIMULAÇÃO (--dry-run): nada foi gravado no banco.'))
        else:
            out.write(self.style.SUCCESS('Importação concluída e gravada no banco.'))

        out.write(f'\nCatálogo anterior substituído: {removed["Exercícios"]} exercícios removidos.')
        out.write('\nRegistros no banco após a importação:')
        for label, count in created.items():
            out.write(f'  {label}: {count}')
        out.write(f'  (exercícios sem equipamento: {rules["free"]})')

        self.write_list(f'Avisos ({len(self.warnings)}):', self.warnings)
        self.write_list(f'Regras do app não atendidas ({len(rules["problems"])}):', rules['problems'])

        out.write('\nPara o time de Educação Física revisar:')
        if rules['review']:
            for status, count in rules['review'].items():
                out.write(f'  - {count} exercício(s) com status "{status or "vazio"}"')
        else:
            out.write('  - todos os exercícios estão com status "Com fonte"')
        for title in rules['flagged_sources']:
            out.write(f'  - fonte marcada "a reconferir": {title}')
        out.write(
            f'  - {len(rules["weak"])} troca(s) "fraca(s)": a alternativa não tem o mesmo grupo principal '
            f'nem o mesmo padrão de movimento' + ('' if options['detalhes'] else ' (use --detalhes para listar)')
        )
        if options['detalhes']:
            for item in rules['weak']:
                out.write(f'      {item}')

    def write_list(self, title, items):
        self.stdout.write(f'\n{title}')
        if not items:
            self.stdout.write('  (nenhum)')
        for item in items:
            self.stdout.write(self.style.WARNING(f'  - {item}'))
